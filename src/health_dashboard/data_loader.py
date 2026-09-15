"""日報エクセルデータの読み込み・クレンジング処理。

就労移行支援事業所の利用者が毎朝入力する日報（Googleフォーム由来のExcel/CSV）から、
入眠時間・起床時間・睡眠の質・中途覚醒回数・気分を抽出し、グラフ化しやすい形に整形する。
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl
import pandas as pd

# 元データの列名（Googleフォームの質問文がそのまま列名になっている）
COL_DATE = "日付"
COL_BEDTIME = "就寝時間"
COL_WAKE = "起床時間"
COL_QUALITY = "睡眠の質"
COL_MOOD_WAKE = "気分（起床時）"
COL_MOOD_COMMUTE = "気分（通所時）"

REQUIRED_COLUMNS = [
    COL_DATE,
    COL_BEDTIME,
    COL_WAKE,
    COL_QUALITY,
    COL_MOOD_WAKE,
    COL_MOOD_COMMUTE,
]

# 睡眠の質ラベル -> スコア（高いほど良い）。ヒアリング通り医学的な重み付けではなく、
# グラフ上で「良好>普通>悪い」の順に並べるための単純な序数。
_QUALITY_SCORE = {"良好": 3, "普通": 2, "悪い": 1}
# スコア -> ラベルの逆引き（グラフの目盛り表示・解説文の生成で共通利用）。
QUALITY_SCORE_LABELS = {score: label for label, score in _QUALITY_SCORE.items()}
# 複数ラベルが併記されていた場合に「より悪い方」を採用するための重大度順。
_QUALITY_SEVERITY = {"良好": 0, "普通": 1, "悪い": 2}

_AWAKENING_RE = re.compile(r"中途覚醒(\d+)(以上)?")

# セルフケアシートのチェックマーク -> ポイント（大きいほど不調）。
SELFCARE_MARK_TO_POINTS = {"〇": 0, "△": 1, "✕": 2}

# 体調ポイントを「悪い」と判定する閾値。
# 実データ（オリジナルセルフケアシート）を集計すると、0点の日を除く非ゼロの日は
# 1〜4点に大半が集中し、7点・13点・20点のように明確な外れ値として離れた日が
# 少数存在する（体調が大きく崩れた日に相当）。この自然な区切りに基づき、
# 5点以上を「悪い」とする。現場の実感と合わない場合はこの値を調整すればよい。
CONDITION_BAD_THRESHOLD = 5

# 表示用の縦軸選択肢: (UI表示名, DataFrameの列名)
AXIS_OPTIONS: list[tuple[str, str]] = [
    ("入眠時間", "bedtime_hours"),
    ("起床時間", "wake_hours"),
    ("睡眠の質", "quality_score"),
    ("中途覚醒回数", "night_awakenings"),
    ("気分（起床時）", "mood_wake"),
    ("気分（通所時）", "mood_commute"),
]


def find_header_row(path: str | Path) -> int:
    """「日付」列を含む行を探し、pandasのheader引数用の0始まり行番号を返す。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if COL_DATE in row:
            return row_idx
    raise ValueError(f"「{COL_DATE}」列を含むヘッダー行が見つかりません: {path}")


def normalize_bedtime(t: dt.time | None) -> dt.time | None:
    """就寝時刻を「夜の時刻」として解釈し直す。

    日報フォームには夜9時台〜11時台の就寝が「09:30」のようにAM/PMの区別なく
    記録されている（例: 就寝9:30・起床6:20で睡眠時間8時間と申告 → 実際は21:30就寝）。
    そのため6時〜11時台の値はPM（+12時間）とみなす。0時〜5時台は「日付をまたいだ後の
    就寝」としてそのまま扱う。
    """
    if t is None:
        return None
    if 6 <= t.hour <= 11:
        return (dt.datetime.combine(dt.date.min, t) + dt.timedelta(hours=12)).time()
    return t


def bedtime_to_chart_hours(t_normalized: dt.time | None) -> float | None:
    """正規化済みの就寝時刻を、深夜またぎでも連続する数値軸に変換する。

    18〜23時台はそのまま18.0〜23.99、0〜5時台（日付またぎ後の就寝）は24を足して
    24.0〜29.99として扱うことで、グラフ上で時系列が途切れないようにする。
    """
    if t_normalized is None:
        return None
    hours = t_normalized.hour + t_normalized.minute / 60
    if t_normalized.hour < 12:
        hours += 24
    return hours


def time_to_hours(t: dt.time | None) -> float | None:
    if t is None:
        return None
    return t.hour + t.minute / 60


def _coerce_time(value: object) -> dt.time | None:
    """openpyxl/pandasが返しうるtime/datetime/文字列表現を統一的にdt.timeへ変換する。"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, dt.datetime):
        return value.time()
    if isinstance(value, dt.time):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return dt.datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return None


def parse_sleep_quality(text: object) -> tuple[str | None, int]:
    """「悪い, 中途覚醒３以上」のような自由記述から (睡眠の質ラベル, 中途覚醒回数) を抽出する。

    - 睡眠の質ラベルが複数併記されている場合は、より悪い方（良好<普通<悪い）を採用する。
    - 「中途覚醒N」「中途覚醒N以上」の数値をそのまま採用する（「以上」は打ち切り値としてNを採用）。
    - 「早朝覚醒」の記載は中途覚醒回数に+1として合算する。
    - 中途覚醒・早朝覚醒のいずれの記載もない場合、中途覚醒回数は0とする。
    """
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return None, 0

    tokens = [tok.strip() for tok in re.split(r"[、,]", str(text)) if tok.strip()]

    quality_label: str | None = None
    for tok in tokens:
        if tok in _QUALITY_SEVERITY:
            if quality_label is None or _QUALITY_SEVERITY[tok] > _QUALITY_SEVERITY[quality_label]:
                quality_label = tok

    awakenings = 0
    for tok in tokens:
        m = _AWAKENING_RE.fullmatch(tok)
        if m:
            awakenings += int(m.group(1))

    if "早朝覚醒" in tokens:
        awakenings += 1

    return quality_label, awakenings


def load_daily_reports(path: str | Path) -> pd.DataFrame:
    """日報Excelを読み込み、分析・グラフ化用に整形したDataFrameを返す。

    戻り値の主な列:
        date, bedtime, bedtime_hours, wake_time, wake_hours,
        quality_label, quality_score, night_awakenings,
        mood_wake, mood_commute
    """
    header_row = find_header_row(path)
    raw = pd.read_excel(path, header=header_row, engine="openpyxl")
    raw = raw.dropna(subset=[COL_DATE]).copy()

    missing = [c for c in REQUIRED_COLUMNS if c not in raw.columns]
    if missing:
        raise ValueError(f"必要な列が見つかりません: {missing}")

    df = pd.DataFrame()
    df["date"] = pd.to_datetime(raw[COL_DATE]).dt.date

    bedtime_raw = raw[COL_BEDTIME].map(_coerce_time)
    bedtime_norm = bedtime_raw.map(normalize_bedtime)
    df["bedtime"] = bedtime_norm
    df["bedtime_hours"] = bedtime_norm.map(bedtime_to_chart_hours)

    wake_time = raw[COL_WAKE].map(_coerce_time)
    df["wake_time"] = wake_time
    df["wake_hours"] = wake_time.map(time_to_hours)

    quality_parsed = raw[COL_QUALITY].map(parse_sleep_quality)
    df["quality_label"] = quality_parsed.map(lambda t: t[0])
    df["quality_score"] = df["quality_label"].map(_QUALITY_SCORE)
    df["night_awakenings"] = quality_parsed.map(lambda t: t[1])

    df["mood_wake"] = pd.to_numeric(raw[COL_MOOD_WAKE], errors="coerce")
    df["mood_commute"] = pd.to_numeric(raw[COL_MOOD_COMMUTE], errors="coerce")

    df = df.sort_values("date").reset_index(drop=True)
    return df


def _find_remarks_column(first_row: tuple) -> int:
    """グループ見出し行（1行目）から「備考」列のインデックスを探す。

    見つからない場合は、テンプレートの既定列数（27列、備考はインデックス26）を仮定する。
    """
    for idx, value in enumerate(first_row):
        if value == "備考":
            return idx
    return 26


def load_selfcare_points(path: str | Path) -> pd.DataFrame:
    """セルフケアシート（月ごとにシートが分かれた〇/△/✕チェック表）から、
    日ごとの体調ポイント（不調ほど高得点）を集計する。

    〇=0点・△=1点・✕=2点として、日付ごとに「睡眠・食事・ストレス」「良好サイン」
    「注意サイン」「悪化サイン」「回復対処」の全チェック項目（備考欄を除く）を合算する。
    シートが複数ある場合はすべて対象にし、同じ日付が複数シートに存在する場合は
    最初に見つかったものを採用する。
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    points_by_date: dict[dt.date, int] = {}
    for ws in wb.worksheets:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        remarks_col = _find_remarks_column(first_row)

        # 1〜4行目は見出し（グループ見出し・項目名・空行・「日付/曜日」）で、
        # 実データは5行目から始まる。
        for row in ws.iter_rows(min_row=5, values_only=True):
            date_value = row[0] if row else None
            if not isinstance(date_value, dt.datetime):
                continue
            marks = row[2:remarks_col]
            if all(m is None for m in marks):
                continue
            points = sum(SELFCARE_MARK_TO_POINTS.get(m, 0) for m in marks if m is not None)
            points_by_date.setdefault(date_value.date(), points)

    if not points_by_date:
        return pd.DataFrame(columns=["date", "condition_points"])

    return (
        pd.DataFrame(
            {"date": list(points_by_date.keys()), "condition_points": list(points_by_date.values())}
        )
        .sort_values("date")
        .reset_index(drop=True)
    )


def condition_label_from_points(points: float | None) -> str | None:
    """体調ポイントを 良好/普通/悪い の3段階ラベルに変換する。

    0点=良好、1点〜(CONDITION_BAD_THRESHOLD - 1)点=普通、
    CONDITION_BAD_THRESHOLD点以上=悪い とする。
    """
    if points is None or (isinstance(points, float) and pd.isna(points)):
        return None
    if points <= 0:
        return "良好"
    if points < CONDITION_BAD_THRESHOLD:
        return "普通"
    return "悪い"


def attach_condition(df: pd.DataFrame, selfcare_df: pd.DataFrame) -> pd.DataFrame:
    """日報データに、セルフケアシート由来の体調ポイント・ラベルを日付でマージする。"""
    merged = df.merge(selfcare_df, on="date", how="left")
    merged["condition_label"] = merged["condition_points"].map(condition_label_from_points)
    return merged


def filter_by_period(df: pd.DataFrame, period: str) -> pd.DataFrame:
    """期間フィルター: '直近1週間' / '直近1ヶ月' / '全期間'。"""
    if df.empty or period == "全期間":
        return df
    latest = max(df["date"])
    if period == "直近1週間":
        start = latest - dt.timedelta(days=6)
    elif period == "直近1ヶ月":
        start = latest - dt.timedelta(days=29)
    else:
        raise ValueError(f"不明な期間指定です: {period}")
    return df[df["date"] >= start].reset_index(drop=True)
