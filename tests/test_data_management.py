import datetime as dt
import os

import openpyxl
import pytest

from health_dashboard.data_loader import find_absence_file, find_user_files, list_user_dirs
from health_dashboard.data_management import (
    ABSENCE,
    DAILY,
    SELFCARE,
    create_user,
    delete_user,
    file_status,
    save_uploaded_file,
    validate_user_name,
)

# --- 利用者の削除 ---


def test_delete_user_removes_folder_and_all_data(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    user_dir, _ = create_user(data_dir, "削除する人")
    other_dir, _ = create_user(data_dir, "残る人")
    assert save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path)) is None
    (user_dir / "memo").mkdir()
    (user_dir / "memo" / "x.txt").write_text("x")
    assert save_uploaded_file(other_dir, DAILY, _daily_bytes(tmp_path)) is None

    assert delete_user(data_dir, user_dir) is None

    assert not user_dir.exists()
    assert [p.name for p in list_user_dirs(data_dir)] == ["残る人"]
    assert find_user_files(other_dir)[0] is not None  # ほかの利用者のデータは残る


def test_delete_user_removes_read_only_files(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    user_dir, _ = create_user(data_dir, "A")
    locked = user_dir / "日報.xlsx"
    locked.write_bytes(b"x")
    locked.chmod(0o444)
    user_dir.chmod(0o755)

    assert delete_user(data_dir, user_dir) is None
    assert not user_dir.exists()


def test_delete_user_refuses_folders_outside_data_dir(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "keep.txt").write_text("x")
    nested = data_dir / "A" / "nested"
    nested.mkdir(parents=True)

    assert delete_user(data_dir, outside) is not None
    assert delete_user(data_dir, data_dir) is not None
    assert delete_user(data_dir, nested) is not None
    assert (outside / "keep.txt").exists()
    assert nested.exists()


def test_delete_user_missing_folder_returns_message(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    assert delete_user(data_dir, data_dir / "いない人") is not None


def test_delete_user_does_not_follow_symlink(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    target = tmp_path / "target"
    target.mkdir()
    (target / "keep.txt").write_text("x")
    link = data_dir / "link"
    try:
        link.symlink_to(target, target_is_directory=True)
    except OSError:
        pytest.skip("シンボリックリンクを作れない環境")

    assert delete_user(data_dir, link) is not None
    assert (target / "keep.txt").exists()

# --- テスト用のExcel（実データと同じ形式の最小限） ---


def _daily_bytes(tmp_path, name="d.xlsx", mood=3.0):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["日付", "就寝時間", "起床時間", "睡眠の質", "気分（起床時）", "気分（通所時）"])
    ws.append([dt.datetime(2026, 6, 4), dt.time(9, 30), dt.time(6, 20), "良好", mood, 6.0])
    wb.save(path)
    return path.read_bytes()


def _selfcare_bytes(tmp_path, name="s.xlsx"):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "202601"
    ws.append([None, None, "グループ", "グループ", "グループ", "備考"])
    ws.append(["チェック項目", None, "A", "B", "C", None])
    ws.append([None] * 6)
    ws.append(["日付", "曜日", None, None, None, None])
    ws.append([dt.datetime(2026, 1, 1), "木", "〇", "△", "✕", None])
    wb.save(path)
    return path.read_bytes()


def _absence_bytes(tmp_path, name="a.xlsx"):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["タイムスタンプ", "欠席理由"])
    ws.append([dt.datetime(2026, 6, 18, 10, 0), "精神不調"])
    wb.save(path)
    return path.read_bytes()


@pytest.fixture
def user_dir(tmp_path):
    folder = tmp_path / "data" / "西村"
    folder.mkdir(parents=True)
    return folder


# --- 利用者の追加 ---


def test_create_user_makes_folder(tmp_path):
    data = tmp_path / "data"

    created, error = create_user(data, "  西村  ")

    assert error is None
    assert created == data / "西村"
    assert created.is_dir()
    assert [p.name for p in list_user_dirs(data)] == ["西村"]


@pytest.mark.parametrize(
    ("name", "message_part"),
    [
        ("", "入力してください"),
        ("   ", "入力してください"),
        ("a/b", "使えない文字"),
        ("山田\\太郎", "使えない文字"),
        ("a:b", "使えない文字"),
        (".hidden", "「.」"),
        ("名前.", "「.」"),
        ("CON", "使えません"),
        ("com1", "使えません"),
        ("あ" * 41, "40文字以内"),
    ],
)
def test_invalid_user_names_are_rejected(name, message_part):
    cleaned, error = validate_user_name(name, [])

    assert cleaned is None
    assert message_part in error


def test_duplicate_user_name_is_rejected_case_insensitively(tmp_path):
    data = tmp_path / "data"
    create_user(data, "Tanaka")

    created, error = create_user(data, "tanaka")

    assert created is None
    assert "すでにいます" in error
    assert len(list_user_dirs(data)) == 1


def test_create_user_does_not_touch_existing_user_files(tmp_path):
    data = tmp_path / "data"
    created, _ = create_user(data, "西村")
    (created / "日報.xlsx").write_text("keep")

    create_user(data, "西村")  # 重複で失敗するだけ

    assert (created / "日報.xlsx").read_text() == "keep"


# --- アップロード ---


def test_upload_saves_each_kind_with_standard_filename(tmp_path, user_dir):
    assert save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path)) is None
    assert save_uploaded_file(user_dir, SELFCARE, _selfcare_bytes(tmp_path)) is None
    assert save_uploaded_file(user_dir, ABSENCE, _absence_bytes(tmp_path)) is None

    assert sorted(p.name for p in user_dir.iterdir()) == [
        "セルフケアシート.xlsx",
        "日報.xlsx",
        "欠席.xlsx",
    ]
    daily, selfcare = find_user_files(user_dir)  # 既存の読み込みの仕組みで見つかる
    assert daily.name == "日報.xlsx" and selfcare.name == "セルフケアシート.xlsx"
    assert find_absence_file(user_dir).name == "欠席.xlsx"


def test_upload_replaces_old_file_of_same_kind_including_other_names(tmp_path, user_dir):
    (user_dir / "西村_日報データ（旧）.xlsx").write_text("old")
    (user_dir / "日報.xlsx").write_text("older")
    (user_dir / "欠席.xlsx").write_text("other kind stays")

    error = save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path, mood=9.0))

    assert error is None
    names = sorted(p.name for p in user_dir.iterdir())
    assert names == ["日報.xlsx", "欠席.xlsx"]  # 古い日報は全て置き換わる
    assert (user_dir / "欠席.xlsx").read_text() == "other kind stays"
    new = openpyxl.load_workbook(user_dir / "日報.xlsx").active
    assert new.cell(row=2, column=5).value == 9.0


def test_wrong_kind_is_rejected_and_old_file_is_kept(tmp_path, user_dir):
    save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path))
    before = (user_dir / "日報.xlsx").read_bytes()

    # 日報の欄に、欠席のファイルを入れてしまった場合
    error = save_uploaded_file(user_dir, DAILY, _absence_bytes(tmp_path))

    assert error is not None
    assert "日報として読み込めませんでした" in error
    assert "今までのデータは、そのまま残しています" in error
    assert (user_dir / "日報.xlsx").read_bytes() == before
    assert [p.name for p in user_dir.iterdir()] == ["日報.xlsx"]  # 一時ファイルも残らない


@pytest.mark.parametrize("kind", [DAILY, SELFCARE, ABSENCE])
def test_corrupt_file_is_rejected_for_every_kind(user_dir, kind):
    error = save_uploaded_file(user_dir, kind, b"this is not an excel file")

    assert error is not None
    assert kind.label in error
    assert list(user_dir.iterdir()) == []


def test_selfcare_without_any_check_is_rejected(tmp_path, user_dir):
    # 日報をセルフケアの欄に入れた場合（チェックの記録が無い）
    error = save_uploaded_file(user_dir, SELFCARE, _daily_bytes(tmp_path))

    assert error is not None
    assert "セルフケアシートとして読み込めませんでした" in error
    assert list(user_dir.iterdir()) == []


def test_absence_file_without_rows_is_accepted(tmp_path, user_dir):
    path = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["タイムスタンプ", "欠席理由"])
    wb.save(path)

    assert save_uploaded_file(user_dir, ABSENCE, path.read_bytes()) is None


def test_locked_old_file_keeps_working_and_reports_clear_message(
    tmp_path, user_dir, monkeypatch
):
    save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path))
    original_unlink = type(user_dir).unlink

    def failing_unlink(self, *args, **kwargs):
        if self.name == "日報.xlsx":
            raise PermissionError(13, "Permission denied")
        return original_unlink(self, *args, **kwargs)

    monkeypatch.setattr(type(user_dir), "unlink", failing_unlink)

    error = save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path, mood=8.0))

    assert error is not None
    assert "Excelで開いている場合は閉じて" in error
    assert not (user_dir / ".upload-tmp.xlsx").exists()
    monkeypatch.undo()
    old = openpyxl.load_workbook(user_dir / "日報.xlsx").active
    assert old.cell(row=2, column=5).value == 3.0  # 元のデータは壊れていない


def test_file_status_reports_registered_files_and_missing_ones(tmp_path, user_dir):
    save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path))

    status = file_status(user_dir)

    path, modified = status["daily"]
    assert path.name == "日報.xlsx"
    assert abs((dt.datetime.now() - modified).total_seconds()) < 60
    assert status["selfcare"] is None
    assert status["absence"] is None


def test_uploaded_data_is_usable_by_the_app_loaders(tmp_path, user_dir):
    from health_dashboard.data_loader import load_daily_reports, load_selfcare_points

    save_uploaded_file(user_dir, DAILY, _daily_bytes(tmp_path))
    save_uploaded_file(user_dir, SELFCARE, _selfcare_bytes(tmp_path))

    daily, selfcare = find_user_files(user_dir)
    assert len(load_daily_reports(daily)) == 1
    assert list(load_selfcare_points(selfcare)["condition_points"]) == [3]


def test_os_independent_temp_name_is_hidden_from_loader(tmp_path, user_dir):
    # 一時ファイルの名前は「日報」等を含まないため、途中で残っても日報として拾われない
    (user_dir / ".upload-tmp.xlsx").write_text("x")

    assert find_user_files(user_dir) == (None, None)
    assert os.path.exists(user_dir / ".upload-tmp.xlsx")
