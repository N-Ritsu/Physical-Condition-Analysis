"""アプリ全体の通しテスト（画面操作）。

利用者の追加 → 日報・セルフケア・欠席のアップロード → 全ての表示・縦軸の切り替えを行い、
どこでも例外が出ないことを確かめる。pandasなどのライブラリの版が変わったときに、
部品単位のテストでは気づけない画面の不具合（例: 欠損が None から NaN に変わる）を見つける。
"""

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

APP_PATH = str(Path(__file__).resolve().parents[1] / "app.py")
DAYS = 45


def _weekdays(count):
    day, found = dt.date(2026, 5, 4), []
    while len(found) < count:
        if day.weekday() < 5:
            found.append(day)
        day += dt.timedelta(days=1)
    return found


def _daily_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["日付", "就寝時間", "起床時間", "睡眠の質"]
    headers += ["気分（起床時）", "気分（通所時）", "気分の種類"]
    ws.append(headers)
    qualities = ["良好", "普通", "悪い, 中途覚醒３以上", "普通, 中途覚醒２", "良好"]
    moods = ["落ち着いている", "不安", "不安, 緊張", "憂鬱", "うれしい"]
    for i, day in enumerate(_weekdays(DAYS)):
        ws.append(
            [
                dt.datetime.combine(day, dt.time()),
                dt.time(9 + (i % 2), 20 * (i % 3)),
                dt.time(6, (i * 7) % 60),
                qualities[i % 5],
                3 + (i * 3) % 7,
                4 + (i * 5) % 6,
                moods[i % 5],
            ]
        )
    wb.save(path)


def _selfcare_file(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("202605")
    ws.append([None, None, "G", "G", "G", "G", "備考"])
    ws.append(["チェック項目", None, "A", "B", "C", "D", None])
    ws.append([None] * 7)
    ws.append(["日付", "曜日", None, None, None, None, None])
    pattern = ["〇", "△", "✕"]
    for i, day in enumerate(_weekdays(DAYS)):
        marks = [pattern[(i + j * (i // 5)) % 3] for j in range(4)]
        ws.append([dt.datetime.combine(day, dt.time()), "月", *marks, None])
    wb.save(path)


def _absence_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["タイムスタンプ", "欠席理由"])
    ws.append([dt.datetime(2026, 5, 20, 10, 0), "精神不調"])
    ws.append([dt.datetime(2026, 6, 3, 10, 0), "体調不良"])
    ws.append([dt.datetime(2026, 6, 10, 10, 0), "通院"])
    wb.save(path)


@pytest.fixture
def app(tmp_path, monkeypatch):
    home = tmp_path / "home"
    (home / "data").mkdir(parents=True)
    monkeypatch.setenv("HEALTH_DASHBOARD_HOME", str(home))
    files = {}
    builders = {"daily": _daily_file, "selfcare": _selfcare_file, "absence": _absence_file}
    for name, builder in builders.items():
        files[name] = tmp_path / f"{name}.xlsx"
        builder(files[name])
    return AppTest.from_file(APP_PATH, default_timeout=180).run(), home, files


def _uploaders(at):
    return {u.label.split("の")[0]: u for u in at.get("file_uploader")}


def _radio(at, label):
    return {r.label: r for r in at.sidebar.radio}[label]


def _assert_no_exception(at, where):
    assert not at.exception, f"{where}: {[e.value for e in at.exception]}"


def test_first_run_shows_add_user_page(app):
    at, _, _ = app

    _assert_no_exception(at, "初回")
    assert at.sidebar.selectbox[0].value == "新しい利用者を追加"
    assert [t.label for t in at.text_input] == ["利用者の名前"]


def test_full_flow_add_user_upload_and_browse_every_page_and_axis(app):
    at, home, files = app

    at.text_input[0].set_value("テスト太郎")
    at.button[0].click().run()
    _assert_no_exception(at, "利用者の追加")
    assert (home / "data" / "テスト太郎").is_dir()
    assert at.sidebar.selectbox[0].value.name == "テスト太郎"

    kinds = (("daily", "日報"), ("selfcare", "セルフケアシート"), ("absence", "欠席情報"))
    for key, label in kinds:
        _uploaders(at)[label].upload(f"{key}.xlsx", files[key].read_bytes()).run()
        _assert_no_exception(at, f"{label}のアップロード")
        assert any(label in s.value for s in at.success), label
    assert sorted(p.name for p in (home / "data" / "テスト太郎").iterdir()) == [
        "セルフケアシート.xlsx",
        "日報.xlsx",
        "欠席.xlsx",
    ]

    # グラフ: 全ての縦軸を、両方の期間で表示する
    _radio(at, "表示").set_value("グラフ").run()
    _assert_no_exception(at, "グラフ")
    assert len(at.get("plotly_chart")) == 1
    for period in ("直近1ヶ月", "全期間"):
        _radio(at, "期間").set_value(period).run()
        for axis in _radio(at, "縦軸（表示する項目）").options:
            _radio(at, "縦軸（表示する項目）").set_value(axis).run()
            _assert_no_exception(at, f"{period} / {axis}")
            assert len(at.get("plotly_chart")) == 1, f"{period} / {axis}"

    for mode in ("出席状況", "相関表"):
        _radio(at, "表示").set_value(mode).run()
        _assert_no_exception(at, mode)

    # 設定のデータ管理 → 表示に戻る
    _radio(at, "設定").set_value("データ管理").run()
    _assert_no_exception(at, "データ管理")
    assert [s.value for s in at.subheader][:1] == ["データ管理"]
    _radio(at, "表示").set_value("グラフ").run()
    _assert_no_exception(at, "データ管理からグラフへ")
    assert len(at.get("plotly_chart")) == 1


def _add_user(at, name):
    at.text_input[0].set_value(name)
    at.button[0].click().run()


def test_weather_location_can_be_set_and_cleared_from_the_app(app):
    at, home, files = app
    _add_user(at, "テスト太郎")
    _uploaders(at)["日報"].upload("daily.xlsx", files["daily"].read_bytes()).run()
    config = home / "weather_config.json"

    _radio(at, "設定").set_value("緯度・経度の設定").run()
    _assert_no_exception(at, "緯度経度の画面")
    assert [t.label for t in at.text_input] == ["緯度（-90 〜 90）", "経度（-180 〜 180）"]
    assert any("未設定" in i.value for i in at.info)

    # 不正な入力は保存しない
    at.text_input[0].set_value("abc")
    at.text_input[1].set_value("200")
    at.button[0].click().run()
    _assert_no_exception(at, "不正な緯度経度")
    assert len(at.error) == 2
    assert not config.exists()

    # 全角でも受け付ける。取得はネットワーク不要になるよう、保存だけを確かめる
    at.text_input[0].set_value("３５．６８")
    at.text_input[1].set_value("139.69")
    at.button[0].click().run()
    _assert_no_exception(at, "緯度経度の保存")
    assert json.loads(config.read_text(encoding="utf-8")) == {
        "latitude": 35.68,
        "longitude": 139.69,
    }
    assert any("緯度 35.68、経度 139.69" in s.value for s in at.success)

    clear = [b for b in at.button if "解除" in b.label][0]
    clear.click().run()
    _assert_no_exception(at, "緯度経度の解除")
    assert json.loads(config.read_text(encoding="utf-8")) == {
        "latitude": None,
        "longitude": None,
    }


def test_user_can_be_deleted_with_their_data(app):
    at, home, files = app
    _add_user(at, "削除する人")
    _uploaders(at)["日報"].upload("daily.xlsx", files["daily"].read_bytes()).run()
    at.sidebar.selectbox[0].set_value("新しい利用者を追加").run()
    _add_user(at, "残る人")
    assert sorted(p.name for p in (home / "data").iterdir()) == ["削除する人", "残る人"]

    user_dirs = {o: i for i, o in enumerate(at.sidebar.selectbox[0].options)}
    at.sidebar.selectbox[0].select_index(user_dirs["削除する人"]).run()
    _radio(at, "設定").set_value("利用者の削除").run()
    _assert_no_exception(at, "削除の画面")

    # 名前が違うと削除しない
    at.text_input[0].set_value("別の名前")
    at.button[0].click().run()
    assert any("一致しません" in e.value for e in at.error)
    assert (home / "data" / "削除する人").exists()

    at.text_input[0].set_value("削除する人")
    at.button[0].click().run()
    _assert_no_exception(at, "削除")
    assert sorted(p.name for p in (home / "data").iterdir()) == ["残る人"]
    assert at.sidebar.selectbox[0].value.name == "残る人"
    assert any("削除しました" in s.value for s in at.success)


def test_last_user_deleted_returns_to_add_user_page(app):
    at, home, _ = app
    _add_user(at, "A")
    # 日報がまだ無い利用者でも、設定から削除できる
    _radio(at, "設定").set_value("利用者の削除").run()
    _assert_no_exception(at, "日報なしの削除画面")

    at.text_input[0].set_value("A")
    at.button[0].click().run()

    _assert_no_exception(at, "最後の利用者の削除")
    assert list((home / "data").iterdir()) == []
    assert at.sidebar.selectbox[0].value == "新しい利用者を追加"
    assert any("削除しました" in s.value for s in at.success)


def test_wrong_file_is_rejected_without_crashing(app):
    at, home, files = app
    at.text_input[0].set_value("テスト太郎")
    at.button[0].click().run()

    _uploaders(at)["日報"].upload("absence.xlsx", files["absence"].read_bytes()).run()

    _assert_no_exception(at, "間違ったファイル")
    assert any("読み込めませんでした" in e.value for e in at.error)
    assert list((home / "data" / "テスト太郎").iterdir()) == []
