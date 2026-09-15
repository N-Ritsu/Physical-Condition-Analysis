import datetime as dt

import openpyxl
import pandas as pd
import pytest

from health_dashboard.data_loader import (
    CONDITION_BAD_THRESHOLD,
    attach_condition,
    condition_label_from_points,
    filter_by_period,
    load_daily_reports,
    load_selfcare_points,
    normalize_bedtime,
    parse_sleep_quality,
)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (dt.time(9, 30), dt.time(21, 30)),
        (dt.time(11, 0), dt.time(23, 0)),
        (dt.time(0, 30), dt.time(0, 30)),
        (dt.time(23, 0), dt.time(23, 0)),
        (None, None),
    ],
)
def test_normalize_bedtime(raw, expected):
    assert normalize_bedtime(raw) == expected


@pytest.mark.parametrize(
    ("text", "expected_label", "expected_awakenings"),
    [
        ("悪い, 中途覚醒３以上", "悪い", 3),
        ("良好", "良好", 0),
        ("普通, 中途覚醒２, 早朝覚醒", "普通", 3),
        ("普通, 悪い, 中途覚醒３以上", "悪い", 3),
        ("良好, 早朝覚醒", "良好", 1),
        (None, None, 0),
    ],
)
def test_parse_sleep_quality(text, expected_label, expected_awakenings):
    label, awakenings = parse_sleep_quality(text)
    assert label == expected_label
    assert awakenings == expected_awakenings


def _write_sample_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "シート1"
    # 実データに合わせ、先頭2行は空、3行目にヘッダーを置く。
    ws.append([None] * 15)
    ws.append([None] * 15)
    ws.append(
        [
            None,
            "日付",
            "曜日",
            "通所時間",
            "目標",
            "就寝時間",
            "起床時間",
            "睡眠時間",
            "睡眠の質",
            "気分（起床時）",
            "気分（通所時）",
        ]
    )
    ws.append(
        [
            None,
            dt.datetime(2026, 5, 28),
            "木",
            dt.time(9, 57),
            "目標A",
            dt.time(9, 30),
            dt.time(6, 50),
            "8時間",
            "悪い, 中途覚醒３以上",
            3.0,
            8.0,
        ]
    )
    ws.append(
        [
            None,
            dt.datetime(2026, 6, 4),
            "木",
            dt.time(10, 0),
            "目標B",
            dt.time(9, 30),
            dt.time(6, 20),
            "8時間",
            "良好",
            7.0,
            6.0,
        ]
    )
    wb.save(path)


def test_load_daily_reports(tmp_path):
    path = tmp_path / "sample.xlsx"
    _write_sample_workbook(path)

    df = load_daily_reports(path)

    assert len(df) == 2
    assert list(df["date"]) == [dt.date(2026, 5, 28), dt.date(2026, 6, 4)]
    # 就寝9:30 -> 21:30として解釈される（12時台以降なのでそのままの数値軸）
    assert df.loc[0, "bedtime_hours"] == pytest.approx(21.5)
    assert df.loc[0, "quality_label"] == "悪い"
    assert df.loc[0, "night_awakenings"] == 3
    assert df.loc[1, "quality_label"] == "良好"
    assert df.loc[1, "night_awakenings"] == 0
    assert df.loc[0, "mood_wake"] == 3.0
    assert df.loc[0, "mood_commute"] == 8.0


def _write_selfcare_workbook(path, sheet_specs):
    """sheet_specs: list of (sheet_title, [(date, weekday, mark_a, mark_b, mark_c, memo), ...])"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheet_specs:
        ws = wb.create_sheet(title)
        ws.append([None, None, "グループ", "グループ", "グループ", "備考"])
        ws.append(["チェック項目", None, "A", "B", "C", None])
        ws.append([None] * 6)
        ws.append(["日付", "曜日", None, None, None, None])
        for row in rows:
            ws.append(list(row))
    wb.save(path)


def test_load_selfcare_points_skips_header_and_blank_rows(tmp_path):
    path = tmp_path / "selfcare.xlsx"
    _write_selfcare_workbook(
        path,
        [
            (
                "202601",
                [
                    (dt.datetime(2026, 1, 1), "木", "〇", "△", "✕", "メモ1"),  # 0+1+2=3
                    (dt.datetime(2026, 1, 2), "金", "〇", "〇", "〇", None),  # 0
                    (dt.datetime(2026, 1, 3), "土", None, None, None, None),  # 記録なし -> 除外
                ],
            )
        ],
    )

    df = load_selfcare_points(path)

    assert list(df["date"]) == [dt.date(2026, 1, 1), dt.date(2026, 1, 2)]
    assert list(df["condition_points"]) == [3, 0]


def test_load_selfcare_points_first_data_row_included(tmp_path):
    """月初(シート5行目=最初のデータ行)が読み飛ばされないことを確認する回帰テスト。"""
    path = tmp_path / "selfcare.xlsx"
    _write_selfcare_workbook(
        path,
        [("202602", [(dt.datetime(2026, 2, 1), "日", "✕", "✕", "✕", None)])],
    )

    df = load_selfcare_points(path)

    assert list(df["date"]) == [dt.date(2026, 2, 1)]
    assert list(df["condition_points"]) == [6]


def test_load_selfcare_points_duplicate_date_keeps_first_sheet(tmp_path):
    path = tmp_path / "selfcare.xlsx"
    _write_selfcare_workbook(
        path,
        [
            ("202601", [(dt.datetime(2026, 1, 1), "木", "〇", "〇", "〇", None)]),  # 0
            ("202601b", [(dt.datetime(2026, 1, 1), "木", "✕", "✕", "✕", None)]),  # 6
        ],
    )

    df = load_selfcare_points(path)

    assert list(df["condition_points"]) == [0]


@pytest.mark.parametrize(
    ("points", "expected"),
    [
        (None, None),
        (float("nan"), None),
        (0, "良好"),
        (1, "普通"),
        (CONDITION_BAD_THRESHOLD - 1, "普通"),
        (CONDITION_BAD_THRESHOLD, "悪い"),
        (CONDITION_BAD_THRESHOLD + 10, "悪い"),
    ],
)
def test_condition_label_from_points(points, expected):
    assert condition_label_from_points(points) == expected


def test_attach_condition():
    df = pd.DataFrame({"date": [dt.date(2026, 1, 1), dt.date(2026, 1, 2)], "value": [1, 2]})
    selfcare_df = pd.DataFrame({"date": [dt.date(2026, 1, 1)], "condition_points": [7]})

    merged = attach_condition(df, selfcare_df)

    assert merged.loc[0, "condition_points"] == 7
    assert merged.loc[0, "condition_label"] == "悪い"
    assert pd.isna(merged.loc[1, "condition_points"])
    assert merged.loc[1, "condition_label"] is None


def test_filter_by_period():
    import pandas as pd

    df = pd.DataFrame(
        {
            "date": [dt.date(2026, 1, 1), dt.date(2026, 1, 15), dt.date(2026, 1, 31)],
            "value": [1, 2, 3],
        }
    )
    result = filter_by_period(df, "直近1週間")
    assert list(result["value"]) == [3]

    result = filter_by_period(df, "全期間")
    assert list(result["value"]) == [1, 2, 3]
